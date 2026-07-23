from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from apps.orders.models import Order, OrderStatus


STATUS_TEXT_MAP = {
    "заказ создан": OrderStatus.IN_PROGRESS,
    "новый": OrderStatus.ACCEPTED,
    "ожидаем ответ/предоплату": OrderStatus.WAITING_RESPONSE,
    "ожидаем ответ от клиента": OrderStatus.WAITING_RESPONSE,
    "позвонили о готовности": OrderStatus.READY,
    "сообщили о готовности": OrderStatus.READY,
    "заказ отдан": OrderStatus.DELIVERED,
    "отменен": OrderStatus.CANCELLED,
    "отменён": OrderStatus.CANCELLED,
}

STATUS_FILL_MAP = {
    "rgb:FFFFFF00": OrderStatus.WAITING_RESPONSE,
    "theme:0:tint:0": OrderStatus.ACCEPTED,
    "theme:0:tint:-0.149876": OrderStatus.READY,
    "theme:6:tint:0": OrderStatus.DELIVERED,
    "rgb:FFFF0000": OrderStatus.CANCELLED,
}

EXCEL_EPOCH = date(1899, 12, 30)


@dataclass
class ImportedOrderRow:
    row_number: int
    order: Order
    warnings: list[str] = field(default_factory=list)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def normalize_status_text(value: Any) -> str:
    return " ".join(clean_text(value).lower().split())


def parse_decimal(value: Any) -> tuple[Decimal, str | None]:
    raw = clean_text(value)
    if not raw:
        return Decimal("0.00"), None

    normalized = raw.replace(" ", "").replace(",", ".")
    try:
        return Decimal(normalized).quantize(Decimal("0.01")), None
    except InvalidOperation:
        return Decimal("0.00"), raw


def normalize_money_fields(
    total_amount: Decimal,
    advance_amount: Decimal,
    additional_payment: Decimal,
) -> tuple[Decimal, Decimal, Decimal, list[str]]:
    warnings: list[str] = []

    if total_amount < 0:
        warnings.append(f"Исходная сумма к оплате была отрицательной: {total_amount}")
        total_amount = Decimal("0.00")
    if advance_amount < 0:
        warnings.append(f"Исходный аванс был отрицательным: {advance_amount}")
        advance_amount = Decimal("0.00")
    if additional_payment < 0:
        warnings.append(f"Исходная доплата была отрицательной: {additional_payment}")
        additional_payment = Decimal("0.00")

    paid_amount = advance_amount + additional_payment
    if paid_amount > total_amount:
        warnings.append(
            "Аванс и доплата больше суммы заказа; сумма к оплате поднята до суммы платежей: "
            f"{advance_amount} + {additional_payment} > {total_amount}"
        )
        total_amount = paid_amount

    return total_amount, advance_amount, additional_payment, warnings


def parse_excel_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return EXCEL_EPOCH + timedelta(days=int(value))

    text = clean_text(value)
    if not text:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return EXCEL_EPOCH + timedelta(days=int(float(text.replace(",", "."))))
    except ValueError:
        return None


def get_fill_key(cell: Cell) -> str:
    color = cell.fill.fgColor
    if color.type == "rgb" and color.rgb:
        rgb = color.rgb.upper()
        if len(rgb) == 8 and rgb.startswith("00"):
            rgb = f"FF{rgb[2:]}"
        return f"rgb:{rgb}"
    if color.type == "theme":
        tint = float(color.tint or 0)
        tint_value = "0" if tint == 0 else f"{tint:.6f}".rstrip("0").rstrip(".")
        return f"theme:{color.theme}:tint:{tint_value}"
    return ""


def get_status(status_cell: Cell, fallback_cell: Cell) -> OrderStatus:
    status_text = normalize_status_text(status_cell.value)
    if status_text in STATUS_TEXT_MAP:
        return STATUS_TEXT_MAP[status_text]

    for cell in (status_cell, fallback_cell):
        fill_key = get_fill_key(cell)
        if fill_key in STATUS_FILL_MAP:
            return STATUS_FILL_MAP[fill_key]

    return OrderStatus.ACCEPTED


def build_comment(
    *,
    extra_values: list[str],
    warnings: list[str],
) -> str:
    lines: list[str] = []
    lines.extend(extra_values)
    if warnings:
        lines.append("Предупреждения импорта:")
        lines.extend(f"- {warning}" for warning in warnings)
    return "\n".join(lines)


def row_has_order_data(values: list[Any]) -> bool:
    meaningful_values = values[4:10] + values[11:]
    return any(clean_text(value) for value in meaningful_values)


def build_order_from_row(
    row: tuple[Cell, ...],
    row_number: int,
    created_by: object,
    default_order_date: date | None,
) -> ImportedOrderRow | None:
    values = [cell.value for cell in row]
    if len(values) < 12:
        values.extend([None] * (12 - len(values)))
    if not row_has_order_data(values):
        return None

    order_date = parse_excel_date(values[0]) or default_order_date or timezone.localdate()
    status = get_status(row[1], row[3])
    delivery_date = parse_excel_date(values[2])
    warnings: list[str] = []
    if status != OrderStatus.DELIVERED:
        if delivery_date is not None:
            warnings.append(
                f"Дата выдачи из Excel не перенесена, потому что статус не 'Заказ отдан': {delivery_date:%d.%m.%Y}"
            )
        delivery_date = None
    elif delivery_date is not None and delivery_date < order_date:
        warnings.append(
            "Дата выдачи из Excel раньше даты заказа и не перенесена: "
            f"{delivery_date:%d.%m.%Y} < {order_date:%d.%m.%Y}"
        )
        delivery_date = None

    order_number = clean_text(values[4])
    work_information = clean_text(values[5]) or "Не указано"
    contacts = clean_text(values[6])

    total_amount, total_warning = parse_decimal(values[7])
    advance_amount, advance_warning = parse_decimal(values[8])
    additional_payment, additional_warning = parse_decimal(values[9])
    total_amount, advance_amount, additional_payment, money_warnings = normalize_money_fields(
        total_amount,
        advance_amount,
        additional_payment,
    )

    warnings.extend(money_warnings)
    for label, raw_value in (
        ("Сумма к оплате", total_warning),
        ("Аванс", advance_warning),
        ("Доплата", additional_warning),
    ):
        if raw_value is not None:
            warnings.append(f"{label} не распознано как число: {raw_value}")

    extra_values = [
        clean_text(value)
        for value in values[11:]
        if clean_text(value)
    ]

    order = Order(
        order_date=order_date,
        delivery_date=delivery_date,
        status=status,
        order_number=order_number,
        work_information=work_information,
        contacts=contacts,
        total_amount=total_amount,
        advance_amount=advance_amount,
        additional_payment=additional_payment,
        comment=build_comment(
            extra_values=extra_values,
            warnings=warnings,
        ),
        created_by=created_by,
    )
    return ImportedOrderRow(row_number=row_number, order=order, warnings=warnings)


class Command(BaseCommand):
    help = "Импортирует заказы из Excel-файла старой таблицы."

    def add_arguments(self, parser) -> None:
        parser.add_argument("xlsx_path", type=str, help="Путь к .xlsx файлу заказов.")
        parser.add_argument(
            "--sheet",
            type=str,
            default=None,
            help="Название листа. По умолчанию используется активный лист.",
        )
        parser.add_argument(
            "--created-by",
            type=str,
            default=None,
            help="Username пользователя, который будет указан создателем заказов.",
        )
        parser.add_argument(
            "--start-row",
            type=int,
            default=5,
            help="Первая строка с заказами. По умолчанию 5.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Только проверить файл и показать сводку, без записи в базу.",
        )
        parser.add_argument(
            "--limit",
            type=int,
            default=None,
            help="Импортировать только первые N подготовленных заказов.",
        )
        parser.add_argument(
            "--last",
            type=int,
            default=None,
            help="Импортировать только последние N подготовленных заказов.",
        )

    def handle(self, *args: object, **options: object) -> None:
        limit = options["limit"]
        last = options["last"]
        if limit is not None and last is not None:
            raise CommandError("Нельзя одновременно использовать --limit и --last.")
        if limit is not None and int(limit) <= 0:
            raise CommandError("--limit должен быть положительным числом.")
        if last is not None and int(last) <= 0:
            raise CommandError("--last должен быть положительным числом.")

        xlsx_path = Path(str(options["xlsx_path"])).expanduser()
        if not xlsx_path.exists():
            raise CommandError(f"Файл не найден: {xlsx_path}")

        created_by = self.get_created_by(options["created_by"])
        workbook = load_workbook(xlsx_path, data_only=True, read_only=False)
        sheet_name = options["sheet"]
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise CommandError(f"Лист не найден: {sheet_name}")
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active

        imported_rows: list[ImportedOrderRow] = []
        skipped_rows = 0
        last_order_date: date | None = None
        start_row = int(options["start_row"])
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=start_row),
            start=start_row,
        ):
            row_date = parse_excel_date(row[0].value)
            if row_date is not None:
                last_order_date = row_date
            imported = build_order_from_row(row, row_number, created_by, last_order_date)
            if imported is None:
                skipped_rows += 1
                continue
            imported_rows.append(imported)

        total_prepared_rows = len(imported_rows)
        if limit is not None:
            imported_rows = imported_rows[: int(limit)]
        if last is not None:
            imported_rows = imported_rows[-int(last):]

        status_counts: dict[str, int] = {}
        warning_count = 0
        for imported in imported_rows:
            status_counts[imported.order.status] = status_counts.get(imported.order.status, 0) + 1
            warning_count += len(imported.warnings)

        self.stdout.write(f"Файл: {xlsx_path}")
        self.stdout.write(f"Лист: {worksheet.title}")
        self.stdout.write(f"Подготовлено заказов всего: {total_prepared_rows}")
        self.stdout.write(f"Выбрано для импорта: {len(imported_rows)}")
        self.stdout.write(f"Пустых строк пропущено: {skipped_rows}")
        self.stdout.write(f"Предупреждений импорта: {warning_count}")
        for status, count in sorted(status_counts.items()):
            self.stdout.write(f"- {OrderStatus(status).label}: {count}")

        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("Dry-run: записи в базу не выполнялись."))
            return

        with transaction.atomic():
            Order.objects.bulk_create([imported.order for imported in imported_rows], batch_size=500)

        self.stdout.write(self.style.SUCCESS(f"Импортировано заказов: {len(imported_rows)}"))

    def get_created_by(self, username: str | None):
        User = get_user_model()
        if username:
            try:
                return User.objects.get(username=username)
            except User.DoesNotExist as exc:
                raise CommandError(f"Пользователь не найден: {username}") from exc

        user = (
            User.objects.filter(is_superuser=True).order_by("pk").first()
            or User.objects.filter(is_staff=True).order_by("pk").first()
            or User.objects.filter(is_active=True).order_by("pk").first()
        )
        if user is None:
            raise CommandError("Нет пользователя для поля created_by. Создайте пользователя или передайте --created-by.")
        return user
