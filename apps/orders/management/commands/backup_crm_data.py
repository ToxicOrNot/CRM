from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path
from typing import Any

from django.conf import settings
from django.core.management.base import BaseCommand
from django.db.models import Prefetch
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from apps.clients.models import Client, ClientContact
from apps.orders.models import Order


HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)


def date_value(value: object) -> str:
    return value.isoformat() if value else ""


def datetime_value(value: object) -> str:
    if not value:
        return ""
    return timezone.localtime(value).isoformat()


def decimal_value(value: Decimal) -> str:
    return str(value.quantize(Decimal("0.01")))


def user_label(user: object | None) -> str:
    if user is None:
        return ""
    full_name = getattr(user, "get_full_name", lambda: "")()
    username = getattr(user, "username", "")
    return full_name or username or str(user)


def client_label(client: Client | None) -> str:
    return str(client) if client else ""


def serialize_contact(contact: ClientContact) -> dict[str, Any]:
    return {
        "id": contact.pk,
        "client_id": contact.client_id,
        "contact_type": contact.contact_type,
        "contact_type_display": contact.get_contact_type_display(),
        "raw_value": contact.raw_value,
        "normalized_value": contact.normalized_value,
        "display_value": contact.display_value,
        "label": contact.label,
        "comment": contact.comment,
        "is_primary": contact.is_primary,
        "created_at": datetime_value(contact.created_at),
        "updated_at": datetime_value(contact.updated_at),
    }


def serialize_client(client: Client) -> dict[str, Any]:
    orders = list(getattr(client, "_prefetched_objects_cache", {}).get("orders", []))
    last_order_date = max((order.order_date for order in orders), default=None)
    contacts = list(getattr(client, "_prefetched_objects_cache", {}).get("contacts", []))
    return {
        "id": client.pk,
        "display_name": client.display_name,
        "preferred_channel": client.preferred_channel,
        "preferred_channel_display": client.get_preferred_channel_display()
        if client.preferred_channel
        else "",
        "primary_contact": client.primary_contact_display,
        "available_channels": client.available_channels_display,
        "comment": client.comment,
        "source_text": client.source_text,
        "created_by_id": client.created_by_id,
        "created_by": user_label(client.created_by),
        "created_at": datetime_value(client.created_at),
        "updated_at": datetime_value(client.updated_at),
        "archived": client.archived,
        "order_count": len(orders),
        "last_order_date": date_value(last_order_date),
        "contacts": [serialize_contact(contact) for contact in contacts],
    }


def serialize_order(order: Order) -> dict[str, Any]:
    return {
        "id": order.pk,
        "display_number": order.display_number,
        "order_date": date_value(order.order_date),
        "delivery_date": date_value(order.delivery_date),
        "status": order.status,
        "status_display": order.get_status_display(),
        "client_id": order.client_id,
        "client": client_label(order.client),
        "order_number": order.order_number,
        "work_information": order.work_information,
        "contacts": order.contacts,
        "original_contacts": order.original_contacts,
        "total_amount": decimal_value(order.total_amount),
        "advance_amount": decimal_value(order.advance_amount),
        "additional_payment": decimal_value(order.additional_payment),
        "balance": decimal_value(order.balance),
        "is_paid": order.is_paid,
        "comment": order.comment,
        "created_by_id": order.created_by_id,
        "created_by": user_label(order.created_by),
        "created_at": datetime_value(order.created_at),
        "updated_at": datetime_value(order.updated_at),
        "archived": order.archived,
    }


def append_rows(sheet: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for row in rows:
        sheet.append(row)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        column_letter = column_cells[0].column_letter
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)


class Command(BaseCommand):
    help = "Create JSON and/or Excel backup files for orders and clients."

    def add_arguments(self, parser) -> None:
        parser.add_argument(
            "--output-dir",
            default=str(settings.BASE_DIR / "backups"),
            help="Directory where backup files will be written.",
        )
        parser.add_argument(
            "--format",
            choices=("json", "xlsx", "both"),
            default="both",
            help="Backup format.",
        )
        parser.add_argument(
            "--prefix",
            default="crm_backup",
            help="Backup file name prefix.",
        )

    def handle(self, *args: object, **options: object) -> None:
        output_dir = Path(str(options["output_dir"]))
        output_dir.mkdir(parents=True, exist_ok=True)

        generated_at = timezone.localtime()
        timestamp = generated_at.strftime("%Y%m%d_%H%M%S")
        prefix = str(options["prefix"])

        orders = list(
            Order.objects.select_related("client", "created_by").order_by("pk"),
        )
        clients = list(
            Client.objects.select_related("created_by")
            .prefetch_related(
                "contacts",
                Prefetch("orders", queryset=Order.objects.order_by("order_date", "pk")),
            )
            .order_by("pk"),
        )
        contacts = list(ClientContact.objects.select_related("client").order_by("client_id", "pk"))

        payload = {
            "generated_at": generated_at.isoformat(),
            "counts": {
                "orders": len(orders),
                "clients": len(clients),
                "client_contacts": len(contacts),
            },
            "orders": [serialize_order(order) for order in orders],
            "clients": [serialize_client(client) for client in clients],
            "client_contacts": [serialize_contact(contact) for contact in contacts],
        }

        created_files: list[Path] = []
        backup_format = str(options["format"])

        if backup_format in {"json", "both"}:
            created_files.append(self.write_json(output_dir, prefix, timestamp, payload))
        if backup_format in {"xlsx", "both"}:
            created_files.append(self.write_xlsx(output_dir, prefix, timestamp, payload))

        for path in created_files:
            self.stdout.write(self.style.SUCCESS(f"Created backup: {path}"))

    def write_json(
        self,
        output_dir: Path,
        prefix: str,
        timestamp: str,
        payload: dict[str, Any],
    ) -> Path:
        path = output_dir / f"{prefix}_{timestamp}.json"
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return path

    def write_xlsx(
        self,
        output_dir: Path,
        prefix: str,
        timestamp: str,
        payload: dict[str, Any],
    ) -> Path:
        path = output_dir / f"{prefix}_{timestamp}.xlsx"
        workbook = Workbook()

        metadata_sheet = workbook.active
        metadata_sheet.title = "Metadata"
        append_rows(
            metadata_sheet,
            ["Поле", "Значение"],
            [
                ["Дата создания backup", payload["generated_at"]],
                ["Заказов", payload["counts"]["orders"]],
                ["Клиентов", payload["counts"]["clients"]],
                ["Контактов клиентов", payload["counts"]["client_contacts"]],
            ],
        )

        orders_sheet = workbook.create_sheet("Orders")
        append_rows(
            orders_sheet,
            [
                "ID",
                "Номер",
                "Дата заказа",
                "Дата выдачи",
                "Статус",
                "Клиент ID",
                "Клиент",
                "Номер заказа",
                "Информация о работе",
                "Контакты",
                "Исходные контакты",
                "Сумма",
                "Аванс",
                "Доплата",
                "Остаток",
                "Оплачен",
                "Комментарий",
                "Создал",
                "Создан",
                "Изменен",
                "Архив",
            ],
            [
                [
                    order["id"],
                    order["display_number"],
                    order["order_date"],
                    order["delivery_date"],
                    order["status_display"],
                    order["client_id"] or "",
                    order["client"],
                    order["order_number"],
                    order["work_information"],
                    order["contacts"],
                    order["original_contacts"],
                    order["total_amount"],
                    order["advance_amount"],
                    order["additional_payment"],
                    order["balance"],
                    order["is_paid"],
                    order["comment"],
                    order["created_by"],
                    order["created_at"],
                    order["updated_at"],
                    order["archived"],
                ]
                for order in payload["orders"]
            ],
        )

        clients_sheet = workbook.create_sheet("Clients")
        append_rows(
            clients_sheet,
            [
                "ID",
                "Имя клиента",
                "Предпочтительный канал",
                "Основной контакт",
                "Доступные каналы",
                "Комментарий",
                "Исходная запись",
                "Создал",
                "Создан",
                "Изменен",
                "Архив",
                "Количество заказов",
                "Дата последнего заказа",
            ],
            [
                [
                    client["id"],
                    client["display_name"],
                    client["preferred_channel_display"],
                    client["primary_contact"],
                    client["available_channels"],
                    client["comment"],
                    client["source_text"],
                    client["created_by"],
                    client["created_at"],
                    client["updated_at"],
                    client["archived"],
                    client["order_count"],
                    client["last_order_date"],
                ]
                for client in payload["clients"]
            ],
        )

        contacts_sheet = workbook.create_sheet("ClientContacts")
        append_rows(
            contacts_sheet,
            [
                "ID",
                "Клиент ID",
                "Тип",
                "Исходное значение",
                "Нормализованное значение",
                "Отображаемое значение",
                "Подпись",
                "Комментарий",
                "Основной",
                "Создан",
                "Изменен",
            ],
            [
                [
                    contact["id"],
                    contact["client_id"],
                    contact["contact_type_display"],
                    contact["raw_value"],
                    contact["normalized_value"],
                    contact["display_value"],
                    contact["label"],
                    contact["comment"],
                    contact["is_primary"],
                    contact["created_at"],
                    contact["updated_at"],
                ]
                for contact in payload["client_contacts"]
            ],
        )

        workbook.save(path)
        return path
