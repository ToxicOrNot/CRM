from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook

from apps.clients.models import Client, ClientContact, ContactType
from apps.orders.models import Order


class BackupCommandTests(TestCase):
    @classmethod
    def setUpTestData(cls) -> None:
        User = get_user_model()
        cls.user = User.objects.create_user(username="backup-user", password="password")
        cls.client_record = Client.objects.create(
            display_name="Backup Client",
            preferred_channel=ContactType.PHONE,
            created_by=cls.user,
        )
        ClientContact.objects.create(
            client=cls.client_record,
            contact_type=ContactType.PHONE,
            raw_value="+7 999 111-22-33",
            is_primary=True,
        )
        cls.order = Order.objects.create(
            client=cls.client_record,
            order_number="B-100",
            work_information="Backup order work",
            contacts="Backup Client +79991112233",
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("20.00"),
            additional_payment=Decimal("5.00"),
            comment="Backup comment",
            created_by=cls.user,
        )

    def test_backup_command_creates_json_with_orders_and_clients(self) -> None:
        with TemporaryDirectory() as temp_dir:
            call_command("backup_crm_data", output_dir=temp_dir, format="json")

            backup_files = list(sorted(Path(temp_dir).glob("*.json")))
            self.assertEqual(len(backup_files), 1)

            payload = json.loads(backup_files[0].read_text(encoding="utf-8"))

        self.assertEqual(payload["counts"]["orders"], 1)
        self.assertEqual(payload["counts"]["clients"], 1)
        self.assertEqual(payload["counts"]["client_contacts"], 1)
        self.assertEqual(payload["orders"][0]["order_number"], "B-100")
        self.assertEqual(payload["orders"][0]["balance"], "75.00")
        self.assertEqual(payload["clients"][0]["display_name"], "Backup Client")
        self.assertEqual(payload["client_contacts"][0]["normalized_value"], "+79991112233")

    def test_backup_command_creates_excel_workbook(self) -> None:
        with TemporaryDirectory() as temp_dir:
            call_command("backup_crm_data", output_dir=temp_dir, format="xlsx")

            backup_files = list(sorted(Path(temp_dir).glob("*.xlsx")))
            self.assertEqual(len(backup_files), 1)
            workbook = load_workbook(backup_files[0])

        self.assertEqual(
            set(workbook.sheetnames),
            {"Metadata", "Orders", "Clients", "ClientContacts"},
        )
        self.assertEqual(workbook["Orders"]["B2"].value, "B-100")
        self.assertEqual(workbook["Orders"]["O2"].value, "75.00")
        self.assertEqual(workbook["Clients"]["B2"].value, "Backup Client")
        self.assertEqual(workbook["ClientContacts"]["E2"].value, "+79991112233")
