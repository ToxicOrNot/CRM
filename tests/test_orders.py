from __future__ import annotations

from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from apps.orders.models import Order, OrderAttachment, OrderStatus
from apps.orders.services.attachment_previews import get_order_attachment_thumbnail_name
from apps.tasks.models import Task, TaskPriority, TaskStatus


class OrderWorkflowTests(TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp_media = TemporaryDirectory()
        cls.media_override = override_settings(MEDIA_ROOT=cls.temp_media.name)
        cls.media_override.enable()
        super().setUpClass()

    @classmethod
    def tearDownClass(cls) -> None:
        super().tearDownClass()
        cls.media_override.disable()
        cls.temp_media.cleanup()

    @classmethod
    def setUpTestData(cls) -> None:
        User = get_user_model()
        cls.user = User.objects.create_user(username="orders-user")
        cls.no_permission_user = User.objects.create_user(username="no-orders")
        cls.task_assignee = User.objects.create_user(username="task-assignee")
        permissions = Permission.objects.filter(
            content_type__app_label="orders",
            codename__in=("add_order", "view_order", "change_order"),
        )
        cls.user.user_permissions.add(*permissions)

    def create_order(self, **kwargs: object) -> Order:
        data = {
            "order_date": timezone.localdate(),
            "status": OrderStatus.ACCEPTED,
            "order_number": "",
            "work_information": "10 фото 10x15 глянец",
            "contacts": "Марина, +79990000000",
            "total_amount": Decimal("1000.00"),
            "advance_amount": Decimal("200.00"),
            "additional_payment": Decimal("100.00"),
            "created_by": self.user,
        }
        data.update(kwargs)
        return Order.objects.create(**data)

    def form_data(self, **overrides: object) -> dict[str, object]:
        data: dict[str, object] = {
            "order_date": timezone.localdate().isoformat(),
            "status": OrderStatus.ACCEPTED,
            "order_number": "ORD-001",
            "work_information": "7 фото 10x15 глянец",
            "contacts": "Клиент, +79990000001",
            "total_amount": "1200.00",
            "advance_amount": "500.00",
            "additional_payment": "200.00",
            "comment": "Комментарий по заказу",
        }
        data.update(overrides)
        return data

    def attach_order_file(self, order: Order, name: str, content: bytes = b"file") -> OrderAttachment:
        return OrderAttachment.objects.create(
            order=order,
            file=SimpleUploadedFile(name, content),
            original_name=name,
            uploaded_by=self.user,
        )

    def test_authenticated_user_can_create_order(self) -> None:
        self.client.force_login(self.user)

        response = self.client.post(reverse("orders:create"), data=self.form_data())

        order = Order.objects.get(order_number="ORD-001")
        self.assertRedirects(response, reverse("orders:list"))

    def test_create_form_does_not_render_delivery_date_or_production_place(self) -> None:
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:create"))

        self.assertNotContains(response, 'name="delivery_date"')
        self.assertNotContains(response, "Место выполнения")

    def test_order_form_has_drag_and_drop_file_upload(self) -> None:
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:create"))

        self.assertContains(response, "data-file-dropzone")
        self.assertContains(response, "data-file-list")
        self.assertContains(response, "js/file_dropzone.js")

    def test_authenticated_user_without_permissions_can_create_order(self) -> None:
        self.client.force_login(self.no_permission_user)

        response = self.client.post(reverse("orders:create"), data=self.form_data())

        self.assertRedirects(response, reverse("orders:list"))
        self.assertTrue(Order.objects.filter(order_number="ORD-001").exists())

    def test_created_by_is_set_to_current_user_on_create(self) -> None:
        self.client.force_login(self.user)

        self.client.post(reverse("orders:create"), data=self.form_data(order_number="ORD-002"))

        order = Order.objects.get(order_number="ORD-002")
        self.assertEqual(order.created_by, self.user)

    def test_balance_is_calculated_correctly(self) -> None:
        order = self.create_order(
            total_amount=Decimal("1000.00"),
            advance_amount=Decimal("300.00"),
            additional_payment=Decimal("125.50"),
        )

        self.assertEqual(order.balance, Decimal("574.50"))

    def test_balance_equals_total_minus_advance_and_additional_payment(self) -> None:
        order = self.create_order(
            total_amount=Decimal("2500.00"),
            advance_amount=Decimal("1000.00"),
            additional_payment=Decimal("250.00"),
        )

        self.assertEqual(
            order.balance,
            order.total_amount - order.advance_amount - order.additional_payment,
        )

    def test_total_amount_cannot_be_negative(self) -> None:
        order = self.create_order()
        order.total_amount = Decimal("-1.00")

        with self.assertRaises(ValidationError):
            order.full_clean()

    def test_advance_amount_cannot_be_negative(self) -> None:
        order = self.create_order()
        order.advance_amount = Decimal("-1.00")

        with self.assertRaises(ValidationError):
            order.full_clean()

    def test_additional_payment_cannot_be_negative(self) -> None:
        order = self.create_order()
        order.additional_payment = Decimal("-1.00")

        with self.assertRaises(ValidationError):
            order.full_clean()

    def test_advance_and_additional_payment_cannot_exceed_total(self) -> None:
        order = self.create_order()
        order.total_amount = Decimal("100.00")
        order.advance_amount = Decimal("70.00")
        order.additional_payment = Decimal("40.00")

        with self.assertRaises(ValidationError):
            order.full_clean()

    def test_delivery_date_cannot_be_before_order_date(self) -> None:
        today = timezone.localdate()
        order = self.create_order(order_date=today, status=OrderStatus.DELIVERED)
        order.delivery_date = today - timedelta(days=1)

        with self.assertRaises(ValidationError):
            order.full_clean()

    def test_overdue_order_is_detected(self) -> None:
        order = self.create_order(
            order_date=timezone.localdate() - timedelta(days=2),
        )
        Order.objects.filter(pk=order.pk).update(
            delivery_date=timezone.localdate() - timedelta(days=1),
        )
        order.refresh_from_db()

        self.assertTrue(order.is_overdue)
        self.assertIn(order, Order.objects.overdue())

    def test_delivered_order_is_not_overdue(self) -> None:
        order = self.create_order(
            order_date=timezone.localdate() - timedelta(days=2),
            status=OrderStatus.DELIVERED,
            delivery_date=timezone.localdate() - timedelta(days=1),
        )

        self.assertFalse(order.is_overdue)

    def test_cancelled_order_is_not_overdue(self) -> None:
        order = self.create_order(
            order_date=timezone.localdate() - timedelta(days=2),
            status=OrderStatus.CANCELLED,
            delivery_date=timezone.localdate() - timedelta(days=1),
        )

        self.assertFalse(order.is_overdue)

    def test_delivery_date_is_set_when_order_is_marked_delivered(self) -> None:
        order = self.create_order(status=OrderStatus.ACCEPTED, delivery_date=None)

        order.status = OrderStatus.DELIVERED
        order.save()

        order.refresh_from_db()
        self.assertEqual(order.delivery_date, timezone.localdate())

    def test_delivery_date_is_cleared_when_order_leaves_delivered_status(self) -> None:
        order = self.create_order(status=OrderStatus.DELIVERED)

        order.status = OrderStatus.IN_PROGRESS
        order.save()

        order.refresh_from_db()
        self.assertIsNone(order.delivery_date)

    def test_archived_order_is_hidden_from_main_list(self) -> None:
        archived_order = self.create_order(order_number="ARCH-1", archived=True)
        active_order = self.create_order(order_number="ACTIVE-1", archived=False)
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        self.assertContains(response, active_order.order_number)
        self.assertNotContains(response, archived_order.order_number)

    def test_archived_order_is_visible_in_archive(self) -> None:
        archived_order = self.create_order(order_number="ARCH-2", archived=True)
        active_order = self.create_order(order_number="ACTIVE-2", archived=False)
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:archive"))

        self.assertContains(response, archived_order.order_number)
        self.assertNotContains(response, active_order.order_number)

    def test_search_by_order_number_works(self) -> None:
        found = self.create_order(order_number="SEARCH-101")
        self.create_order(order_number="OTHER-101")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"), {"q": "search-101"})

        self.assertContains(response, found.order_number)
        self.assertNotContains(response, "OTHER-101")

    def test_order_list_partial_returns_rows_for_infinite_scroll(self) -> None:
        for index in range(30):
            self.create_order(order_number=f"INF-{index:02d}")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"), {"page": "2", "partial": "1"})

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "orders/_order_rows.html")
        self.assertContains(response, "order-summary-row")
        self.assertNotContains(response, "<h1")

    def test_order_list_marks_orders_with_comments(self) -> None:
        self.create_order(order_number="COMMENT-1", comment="Позвонить перед выдачей")
        self.create_order(order_number="NO-COMMENT")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        self.assertContains(response, "Есть комментарий")
        self.assertContains(response, "order-comment-indicator")

    def test_authenticated_user_can_attach_file_on_create(self) -> None:
        self.client.force_login(self.user)
        uploaded_file = SimpleUploadedFile(
            "order-brief.txt",
            b"Order file content",
            content_type="text/plain",
        )

        response = self.client.post(
            reverse("orders:create"),
            data={**self.form_data(order_number="FILE-CREATE-1"), "attachments": uploaded_file},
        )

        self.assertRedirects(response, reverse("orders:list"))
        order = Order.objects.get(order_number="FILE-CREATE-1")
        attachment = OrderAttachment.objects.get(order=order)
        self.assertEqual(attachment.original_name, "order-brief.txt")
        self.assertEqual(attachment.uploaded_by, self.user)
        self.assertTrue(attachment.file.name.startswith(f"orders/{order.pk}/attachments/"))

    def test_authenticated_user_can_attach_more_than_one_hundred_files(self) -> None:
        self.client.force_login(self.user)
        uploaded_files = [
            SimpleUploadedFile(
                f"photo-{index:03d}.txt",
                b"content",
                content_type="text/plain",
            )
            for index in range(183)
        ]

        response = self.client.post(
            reverse("orders:create"),
            data={**self.form_data(order_number="MANY-FILES-1"), "attachments": uploaded_files},
        )

        self.assertRedirects(response, reverse("orders:list"))
        order = Order.objects.get(order_number="MANY-FILES-1")
        self.assertEqual(order.attachments.count(), 183)

    def test_order_list_marks_orders_with_attachments(self) -> None:
        order = self.create_order(order_number="FILE-MARK-1")
        OrderAttachment.objects.create(
            order=order,
            file=SimpleUploadedFile("document.txt", b"content"),
            original_name="document.txt",
            uploaded_by=self.user,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        self.assertContains(response, "Есть файлы")
        self.assertContains(response, "order-file-indicator")
        self.assertContains(response, "plus.png")
        self.assertContains(response, "background-color: transparent")
        self.assertContains(response, "icon.png")
        self.assertContains(response, "document.txt")

    def test_order_attachment_detects_image_by_extension(self) -> None:
        order = self.create_order(order_number="IMAGE-FILE-1")

        image_attachment = self.attach_order_file(order, "photo.JPG")
        document_attachment = self.attach_order_file(order, "document.pdf")

        self.assertTrue(image_attachment.is_image)
        self.assertFalse(document_attachment.is_image)

    def test_heic_attachment_uses_preview_endpoint(self) -> None:
        order = self.create_order(order_number="HEIC-PREVIEW-1")
        attachment = self.attach_order_file(order, "iphone-photo.HEIC")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        preview_url = reverse("orders:attachment_preview", kwargs={"pk": attachment.pk})
        self.assertTrue(attachment.is_image)
        self.assertTrue(attachment.is_heic_image)
        self.assertEqual(attachment.preview_url, preview_url)
        self.assertContains(response, f'src="{preview_url}"')
        self.assertContains(response, f'href="{attachment.file.url}"')

    def test_jpeg_attachment_uses_preview_endpoint(self) -> None:
        order = self.create_order(order_number="JPG-PREVIEW-1")
        attachment = self.attach_order_file(order, "camera-photo.jpg")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        preview_url = reverse("orders:attachment_preview", kwargs={"pk": attachment.pk})
        self.assertTrue(attachment.is_image)
        self.assertFalse(attachment.is_heic_image)
        self.assertEqual(attachment.preview_url, preview_url)
        self.assertContains(response, f'src="{preview_url}"')
        self.assertContains(response, f'href="{attachment.file.url}"')

    def test_heic_preview_endpoint_returns_jpeg(self) -> None:
        from PIL import Image
        from pillow_heif import register_heif_opener

        register_heif_opener()
        image = Image.new("RGB", (8, 8), "red")
        heic_file = BytesIO()
        image.save(heic_file, format="HEIF")
        order = self.create_order(order_number="HEIC-CONVERT-1")
        attachment = self.attach_order_file(order, "preview.heic", heic_file.getvalue())
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:attachment_preview", kwargs={"pk": attachment.pk}))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "image/jpeg")
        self.assertTrue(response.content.startswith(b"\xff\xd8"))

    def test_regular_image_preview_endpoint_returns_compressed_thumbnail(self) -> None:
        from PIL import Image

        image = Image.new("RGB", (1400, 1000), "blue")
        image_file = BytesIO()
        image.save(image_file, format="JPEG", quality=95)
        original_content = image_file.getvalue()
        order = self.create_order(order_number="JPG-THUMB-1")
        attachment = self.attach_order_file(order, "large-photo.jpg", original_content)
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:attachment_preview", kwargs={"pk": attachment.pk}))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "image/jpeg")
        self.assertTrue(response.content.startswith(b"\xff\xd8"))
        with Image.open(BytesIO(response.content)) as thumbnail:
            self.assertLessEqual(max(thumbnail.size), 360)
        self.assertTrue(attachment.file.storage.exists(get_order_attachment_thumbnail_name(attachment)))

    def test_uploaded_image_creates_cached_preview(self) -> None:
        from PIL import Image

        self.client.force_login(self.user)
        image = Image.new("RGB", (1200, 800), "green")
        image_file = BytesIO()
        image.save(image_file, format="JPEG", quality=95)
        uploaded_file = SimpleUploadedFile(
            "uploaded-photo.jpg",
            image_file.getvalue(),
            content_type="image/jpeg",
        )

        response = self.client.post(
            reverse("orders:create"),
            data={**self.form_data(order_number="JPG-UPLOAD-1"), "attachments": uploaded_file},
        )

        self.assertRedirects(response, reverse("orders:list"))
        attachment = OrderAttachment.objects.get(order__order_number="JPG-UPLOAD-1")
        self.assertTrue(attachment.file.storage.exists(get_order_attachment_thumbnail_name(attachment)))

    def test_order_list_shows_limited_photo_preview_and_extra_count(self) -> None:
        order = self.create_order(order_number="IMAGE-PREVIEW-1")
        for index in range(1, 43):
            self.attach_order_file(order, f"photo-{index}.jpg")
        self.attach_order_file(order, "contract.pdf")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        self.assertContains(response, 'class="order-photo-thumb"', count=39)
        self.assertContains(response, 'class="order-photo-caption"', count=39)
        self.assertContains(response, "order-photo-more")
        self.assertContains(response, "+3")
        self.assertContains(response, "Файлы (43)")
        self.assertNotContains(response, 'alt="photo-')

    def test_order_detail_shows_attachment_count_near_files_heading(self) -> None:
        order = self.create_order(order_number="FILE-COUNT-1")
        self.attach_order_file(order, "first.jpg")
        self.attach_order_file(order, "second.txt")
        self.client.force_login(self.user)

        response = self.client.get(order.get_absolute_url())

        self.assertContains(response, "Файлы (2)")

    def test_order_edit_form_shows_delete_button_for_existing_attachment(self) -> None:
        order = self.create_order(order_number="FILE-DELETE-FORM-1")
        attachment = self.attach_order_file(order, "wrong-photo.jpg")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:edit", kwargs={"pk": order.pk}))

        self.assertContains(response, "wrong-photo.jpg")
        self.assertContains(response, reverse("orders:delete_attachment", kwargs={"pk": attachment.pk}))
        self.assertContains(response, "Удалить")

    def test_order_edit_form_does_not_nest_delete_form_inside_order_form(self) -> None:
        order = self.create_order(order_number="FILE-NESTED-FORM-1")
        attachment = self.attach_order_file(order, "wrong-photo.jpg")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:edit", kwargs={"pk": order.pk}))

        content = response.content.decode()
        order_form_start = content.index('id="order-form"')
        order_form_end = content.index("</form>", order_form_start)
        delete_form_position = content.index(
            reverse("orders:delete_attachment", kwargs={"pk": attachment.pk}),
        )
        self.assertGreater(delete_form_position, order_form_end)

    def test_order_edit_form_submission_updates_order_with_existing_attachment(self) -> None:
        order = self.create_order(order_number="EDIT-BEFORE")
        self.attach_order_file(order, "existing-photo.jpg")
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("orders:edit", kwargs={"pk": order.pk}),
            data=self.form_data(order_number="EDIT-AFTER", work_information="Новая информация"),
        )

        order.refresh_from_db()
        self.assertRedirects(response, order.get_absolute_url())
        self.assertEqual(order.order_number, "EDIT-AFTER")
        self.assertEqual(order.work_information, "Новая информация")

    def test_authenticated_user_can_delete_order_attachment(self) -> None:
        order = self.create_order(order_number="FILE-DELETE-1")
        attachment = self.attach_order_file(order, "remove-me.jpg")
        file_name = attachment.file.name
        storage = attachment.file.storage
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("orders:delete_attachment", kwargs={"pk": attachment.pk}),
            data={"next": order.get_absolute_url()},
        )

        self.assertRedirects(response, order.get_absolute_url())
        self.assertFalse(OrderAttachment.objects.filter(pk=attachment.pk).exists())
        self.assertFalse(storage.exists(file_name))

    def test_deleting_order_attachment_removes_cached_preview(self) -> None:
        from PIL import Image

        image = Image.new("RGB", (1200, 800), "red")
        image_file = BytesIO()
        image.save(image_file, format="JPEG", quality=95)
        order = self.create_order(order_number="FILE-DELETE-THUMB-1")
        attachment = self.attach_order_file(order, "remove-preview.jpg", image_file.getvalue())
        storage = attachment.file.storage
        thumbnail_name = get_order_attachment_thumbnail_name(attachment)
        self.client.force_login(self.user)

        self.client.get(reverse("orders:attachment_preview", kwargs={"pk": attachment.pk}))
        self.assertTrue(storage.exists(thumbnail_name))

        response = self.client.post(
            reverse("orders:delete_attachment", kwargs={"pk": attachment.pk}),
            data={"next": order.get_absolute_url()},
        )

        self.assertRedirects(response, order.get_absolute_url())
        self.assertFalse(storage.exists(thumbnail_name))

    def test_can_download_all_order_attachments_as_zip(self) -> None:
        order = self.create_order(order_number="FILE-ZIP-1")
        OrderAttachment.objects.create(
            order=order,
            file=SimpleUploadedFile("first.txt", b"first"),
            original_name="first.txt",
            uploaded_by=self.user,
        )
        OrderAttachment.objects.create(
            order=order,
            file=SimpleUploadedFile("second.txt", b"second"),
            original_name="second.txt",
            uploaded_by=self.user,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:download_attachments", kwargs={"pk": order.pk}))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        with ZipFile(BytesIO(response.content)) as zip_file:
            self.assertEqual(set(zip_file.namelist()), {"first.txt", "second.txt"})
            self.assertEqual(zip_file.read("first.txt"), b"first")

    def test_search_by_contacts_works(self) -> None:
        found = self.create_order(order_number="CONTACT-1", contacts="telegram @client_one")
        self.create_order(order_number="CONTACT-2", contacts="phone +7999")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"), {"q": "@client_one"})

        self.assertContains(response, found.order_number)
        self.assertNotContains(response, "CONTACT-2")

    def test_phone_search_ignores_separators_in_query(self) -> None:
        found = self.create_order(order_number="PHONE-PLAIN", contacts="89854337308")
        self.create_order(order_number="PHONE-OTHER", contacts="89111111111")
        self.client.force_login(self.user)

        for query in ("73-08", "73 08", "7308"):
            with self.subTest(query=query):
                response = self.client.get(reverse("orders:list"), {"q": query})
                self.assertContains(response, found.order_number)
                self.assertNotContains(response, "PHONE-OTHER")

    def test_phone_search_ignores_separators_in_contacts(self) -> None:
        found = self.create_order(order_number="PHONE-FORMATTED", contacts="+7 985 433-73-08")
        self.create_order(order_number="PHONE-OTHER", contacts="8 911 111-11-11")
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"), {"q": "7308"})

        self.assertContains(response, found.order_number)
        self.assertNotContains(response, "PHONE-OTHER")

    def test_status_filter_works(self) -> None:
        ready = self.create_order(order_number="READY-1", status=OrderStatus.READY)
        self.create_order(order_number="ACCEPTED-1", status=OrderStatus.ACCEPTED)
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"), {"status": OrderStatus.READY})

        self.assertContains(response, ready.order_number)
        self.assertNotContains(response, "ACCEPTED-1")

    def test_unpaid_filter_works(self) -> None:
        unpaid = self.create_order(
            order_number="DUE-1",
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("50.00"),
            additional_payment=Decimal("0.00"),
        )
        self.create_order(
            order_number="SETTLED-1",
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("50.00"),
            additional_payment=Decimal("50.00"),
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"), {"unpaid": "1"})

        self.assertContains(response, unpaid.order_number)
        self.assertNotContains(response, "SETTLED-1")

    def test_paid_order_has_badge_without_paid_row_color(self) -> None:
        self.create_order(
            order_number="PAID-BADGE-1",
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("40.00"),
            additional_payment=Decimal("60.00"),
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        self.assertContains(response, "Оплачено")
        self.assertNotContains(response, "table-success")

    def test_order_list_highlights_balance_by_payment_state(self) -> None:
        self.create_order(
            order_number="BALANCE-DUE-1",
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("40.00"),
            additional_payment=Decimal("10.00"),
        )
        self.create_order(
            order_number="BALANCE-PAID-1",
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("40.00"),
            additional_payment=Decimal("60.00"),
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        self.assertContains(response, 'class="order-balance-value order-balance-due"')
        self.assertContains(response, 'class="order-balance-value order-balance-paid"')
        self.assertContains(response, 'class="badge text-bg-danger">Остаток 50,00</span>')

    def test_order_detail_highlights_balance_by_payment_state(self) -> None:
        due_order = self.create_order(
            order_number="BALANCE-DUE-DETAIL-1",
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("25.00"),
            additional_payment=Decimal("25.00"),
        )
        paid_order = self.create_order(
            order_number="BALANCE-PAID-DETAIL-1",
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("50.00"),
            additional_payment=Decimal("50.00"),
        )
        self.client.force_login(self.user)

        due_response = self.client.get(due_order.get_absolute_url())
        paid_response = self.client.get(paid_order.get_absolute_url())

        self.assertContains(due_response, 'class="order-balance-value order-balance-due"')
        self.assertContains(paid_response, 'class="order-balance-value order-balance-paid"')

    def test_order_statuses_have_expected_row_colors(self) -> None:
        self.create_order(order_number="WAIT-1", status=OrderStatus.WAITING_RESPONSE)
        self.create_order(order_number="WORK-1", status=OrderStatus.IN_PROGRESS)
        self.create_order(order_number="READY-CALL-1", status=OrderStatus.READY)
        self.create_order(order_number="DONE-1", status=OrderStatus.DELIVERED)
        self.create_order(order_number="CANCEL-1", status=OrderStatus.CANCELLED)
        self.client.force_login(self.user)

        response = self.client.get(reverse("orders:list"))

        self.assertContains(response, "order-row-waiting")
        self.assertContains(response, "order-row-in-progress")
        self.assertContains(response, "order-row-ready")
        self.assertContains(response, "order-row-delivered")
        self.assertContains(response, "order-row-cancelled")

    def test_quick_update_changes_status_and_payment_amounts(self) -> None:
        order = self.create_order(
            order_number="QUICK-1",
            status=OrderStatus.ACCEPTED,
            total_amount=Decimal("100.00"),
            advance_amount=Decimal("10.00"),
            additional_payment=Decimal("0.00"),
            delivery_date=None,
        )
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("orders:quick_update", kwargs={"pk": order.pk}),
            data={
                f"order-{order.pk}-status": OrderStatus.DELIVERED,
                f"order-{order.pk}-total_amount": "100.00",
                f"order-{order.pk}-advance_amount": "40.00",
                f"order-{order.pk}-additional_payment": "60.00",
                "next": reverse("orders:list"),
            },
        )

        self.assertRedirects(response, reverse("orders:list"))
        order.refresh_from_db()
        self.assertEqual(order.status, OrderStatus.DELIVERED)
        self.assertEqual(order.advance_amount, Decimal("40.00"))
        self.assertEqual(order.additional_payment, Decimal("60.00"))
        self.assertEqual(order.balance, Decimal("0.00"))
        self.assertEqual(order.delivery_date, timezone.localdate())

    def test_linked_task_is_visible_on_order_detail(self) -> None:
        order = self.create_order(order_number="TASK-ORDER-1")
        task = Task.objects.create(
            title="Задача по заказу",
            creator=self.user,
            assignee=self.task_assignee,
            order=order,
            priority=TaskPriority.NORMAL,
            status=TaskStatus.NEW,
        )
        self.client.force_login(self.user)

        response = self.client.get(order.get_absolute_url())

        self.assertContains(response, task.title)
        self.assertContains(response, reverse("tasks:detail", kwargs={"pk": task.pk}))

    def test_archiving_order_does_not_delete_linked_tasks(self) -> None:
        order = self.create_order(order_number="ARCHIVE-WITH-TASK")
        task = Task.objects.create(
            title="Не удалять",
            creator=self.user,
            assignee=self.task_assignee,
            order=order,
            priority=TaskPriority.NORMAL,
            status=TaskStatus.NEW,
        )
        self.client.force_login(self.user)

        response = self.client.post(reverse("orders:archive_item", kwargs={"pk": order.pk}))

        self.assertRedirects(response, reverse("orders:archive"))
        order.refresh_from_db()
        self.assertTrue(order.archived)
        self.assertTrue(Task.objects.filter(pk=task.pk, order=order).exists())

    def test_comment_is_saved_and_visible_on_detail(self) -> None:
        order = self.create_order(
            order_number="COMMENT-1",
            comment="Срочно напечатать до вечера",
        )
        self.client.force_login(self.user)

        response = self.client.get(order.get_absolute_url())

        self.assertContains(response, "Срочно напечатать до вечера")

    def test_order_number_may_be_blank(self) -> None:
        first = self.create_order(order_number="")
        second = self.create_order(order_number="")

        self.assertEqual(first.order_number, "")
        self.assertEqual(second.order_number, "")

    def test_display_number_uses_sequence_number_when_order_number_is_blank(self) -> None:
        order = self.create_order(order_number="")

        self.assertEqual(order.sequence_number, str(order.pk))
        self.assertEqual(order.display_number, str(order.pk))
        self.assertEqual(str(order), f"Заказ {order.pk}")

    def test_display_number_prefers_order_number_when_it_exists(self) -> None:
        order = self.create_order(order_number="ORD-777")

        self.assertEqual(order.display_number, "ORD-777")
        self.assertEqual(str(order), "Заказ ORD-777")

    def test_duplicate_non_empty_active_order_number_is_allowed(self) -> None:
        self.create_order(order_number="DUP-1")

        duplicate = self.create_order(order_number=" dup-1 ")

        self.assertEqual(duplicate.order_number, "dup-1")

    def test_import_orders_xlsx_allows_duplicate_order_numbers(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "orders"
        worksheet.append([])
        worksheet.append([])
        worksheet.append([])
        worksheet.append([
            "дата",
            "статус",
            "дата выдачи",
            "где печатаем",
            "номер заказа",
            "Информация о работе",
            "контакты",
            "сумма к оплате",
            "Аванс",
            "Доплата",
            "Остаток",
            "Комментарий",
        ])
        worksheet.append([
            timezone.localdate(),
            "Заказ создан",
            "",
            "Лаба",
            "DUP-XLSX",
            "Первый заказ",
            "89854337308",
            100,
            50,
            50,
            0,
            "первая строка",
        ])
        worksheet.append([
            timezone.localdate(),
            "",
            "",
            "Лаба",
            "DUP-XLSX",
            "Второй заказ",
            "73-08",
            200,
            "",
            "",
            200,
            "",
        ])
        worksheet.append([
            timezone.localdate(),
            "",
            "",
            "Лаба",
            "OVERPAID-XLSX",
            "Overpaid order",
            "Contact",
            100,
            150,
            0,
            -50,
            "",
        ])
        worksheet["D6"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")

        with TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir) / "orders.xlsx"
            workbook.save(temp_path)

            call_command(
                "import_orders_xlsx",
                temp_path,
                "--created-by",
                self.user.username,
                stdout=StringIO(),
            )

        imported_orders = Order.objects.filter(order_number="DUP-XLSX").order_by("pk")
        self.assertEqual(imported_orders.count(), 2)
        self.assertEqual(imported_orders[0].status, OrderStatus.IN_PROGRESS)
        self.assertEqual(imported_orders[1].status, OrderStatus.WAITING_RESPONSE)
        self.assertNotIn("Где печатаем:", imported_orders[0].comment)
        self.assertIn("первая строка", imported_orders[0].comment)
        overpaid_order = Order.objects.get(order_number="OVERPAID-XLSX")
        self.assertEqual(overpaid_order.total_amount, Decimal("150.00"))
        self.assertEqual(overpaid_order.balance, Decimal("0.00"))
        self.assertIn("150.00 + 0.00 > 100.00", overpaid_order.comment)

    def test_import_orders_xlsx_can_import_last_rows_only(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "orders"
        worksheet.append([])
        worksheet.append([])
        worksheet.append([])
        worksheet.append(["date", "status", "delivery", "place", "number", "work", "contacts", "total"])
        for index in range(1, 4):
            worksheet.append([
                timezone.localdate(),
                "",
                "",
                "",
                f"LAST-{index}",
                f"Order {index}",
                f"Contact {index}",
                100,
            ])

        with TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir) / "orders.xlsx"
            workbook.save(temp_path)

            call_command(
                "import_orders_xlsx",
                temp_path,
                "--created-by",
                self.user.username,
                "--last",
                "2",
                stdout=StringIO(),
            )

        self.assertFalse(Order.objects.filter(order_number="LAST-1").exists())
        self.assertTrue(Order.objects.filter(order_number="LAST-2").exists())
        self.assertTrue(Order.objects.filter(order_number="LAST-3").exists())
